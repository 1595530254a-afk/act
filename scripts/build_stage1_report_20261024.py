"""Build the evidence-based stage1 Word base and portable video companions.

This assembles saved results only. It performs no training, physics or robot IO.
Use the document runtime's Python; finalize embedded MP4 packages with the PS tool.
"""
import hashlib
import json
from pathlib import Path
import shutil
import subprocess
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'scripts'))
from cpu_budget import CPU_BUDGET

import matplotlib
matplotlib.use('Agg')
from matplotlib import pyplot as plt
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = ROOT / 'output/ACT第一阶段记录_20261024'
WORK = ROOT / 'tmp/stage1_report_20261024'
EVIDENCE = ROOT / 'reports/stage1_archive_20261024/history_evidence.json'
H = json.loads(EVIDENCE.read_text(encoding='utf-8'))
COMMIT = subprocess.check_output(['git', 'rev-parse', 'HEAD'], cwd=ROOT, text=True).strip()
FONT = 'Microsoft YaHei'
plt.rcParams.update({'font.sans-serif': [FONT], 'axes.unicode_minus': False, 'font.size': 10})


def sha(path):
    with Path(path).open('rb') as f:
        return hashlib.file_digest(f, 'sha256').hexdigest()


def write_json(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')


def set_font(run, size=10.5, bold=False, color='000000'):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), FONT)


def para(text='', size=10.5, bold=False, after=5):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.2
    set_font(p.add_run(text), size, bold)
    return p


def heading(text, level=1):
    return doc.add_heading(text, level=level)


def page(title):
    p = heading(title)
    p.paragraph_format.page_break_before = True


def sources(text):
    para('证据索引  ' + text, 8, False, 2)


def hyperlink(p, label, target):
    rel = p.part.relate_to(target, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    h = OxmlElement('w:hyperlink'); h.set(qn('r:id'), rel)
    r = OxmlElement('w:r'); rp = OxmlElement('w:rPr')
    c = OxmlElement('w:color'); c.set(qn('w:val'), '245C88'); rp.append(c)
    u = OxmlElement('w:u'); u.set(qn('w:val'), 'single'); rp.append(u)
    fonts = OxmlElement('w:rFonts'); fonts.set(qn('w:eastAsia'), FONT); rp.append(fonts)
    sz = OxmlElement('w:sz'); sz.set(qn('w:val'), '18'); rp.append(sz)
    r.append(rp); t = OxmlElement('w:t'); t.text = label; r.append(t); h.append(r); p._p.append(h)


def table(headers, rows, widths, size=9):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    for i, w in enumerate(widths):
        t.columns[i].width = Cm(w)
    for i, s in enumerate(headers):
        t.rows[0].cells[i].text = s
    for row in rows:
        for cell, value in zip(t.add_row().cells, row):
            cell.text = str(value)
    for n, row in enumerate(t.rows):
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        if n == 0:
            row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        for i, cell in enumerate(row.cells):
            cell.width = Cm(widths[i])
            tcpr = cell._tc.get_or_add_tcPr()
            mar = OxmlElement('w:tcMar')
            for side in ('top', 'bottom', 'left', 'right'):
                x = OxmlElement('w:' + side); x.set(qn('w:w'), '75'); x.set(qn('w:type'), 'dxa'); mar.append(x)
            tcpr.append(mar)
            if n == 0:
                sh = OxmlElement('w:shd'); sh.set(qn('w:fill'), 'EEEEEE'); tcpr.append(sh)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.1
                for r in p.runs:
                    set_font(r, size, n == 0)
    borders = OxmlElement('w:tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        b = OxmlElement('w:' + side); b.set(qn('w:val'), 'single'); b.set(qn('w:sz'), '4'); b.set(qn('w:color'), 'CCCCCC'); borders.append(b)
    t._tbl.tblPr.append(borders)
    para('', after=1)
    return t


def picture(path, width=15.8, caption=''):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(path), width=Cm(width))
    if caption:
        c = para(caption, 8.5); c.alignment = WD_ALIGN_PARAGRAPH.CENTER


def video(vid, description):
    item = next(e for e in media if e['id'] == vid)
    title = para(vid + '  ' + description, 10, True, 2)
    title.paragraph_format.keep_with_next = True
    p = para('[[VIDEO:' + vid + ']]', 9, False, 2)
    p.paragraph_format.keep_with_next = True
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(5)
    hyperlink(p, '外部备份 ' + vid + '.mp4', '视频附件/' + vid + '.mp4')


def build_media():
    names = [
        ('V01', 'ACT_平顶红块抓取入盒_仿真成功_20261023.mp4', '单位置sim300成功'),
        ('V02', 'ACT_位置泛化_六组同步回看_20261023.mp4', '固定示范后换位置失败'),
        ('V03', 'ACT_多位置_final_heldout_v1_step600_同步回看.mp4', '旧多位置step600六点测试'),
        ('V04', 'ACT_5000_final8_selected_v1_ensemble1_6ef257e2_同步回看.mp4', '最终八点同步回看'),
    ]
    for i in range(8):
        names.append((f'V{i+5:02d}', f'ACT_5000_final8_selected_v1_ensemble1_6ef257e2_newtest_{i:02d}_实际输入.mp4', f'最终测试newtest_{i:02d}实际双路输入'))
    names.extend([
        ('V13', 'ACT_5000_extra_p10_0_rollout_ensemble1_0d81cc98_extra_p10_0_实际输入.mp4', '追加位置(+10,0)mm实际输入'),
        ('V14', 'ACT_ALOHA_官方基线_seed20261024_20261024.mp4', '独立ALOHA基线一回合'),
    ])
    entries = []
    for vid, name, label in names:
        src = ROOT / '视频' / name
        dst = OUT / '视频附件' / (vid + '.mp4')
        dst.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(src, dst)
        assert sha(src) == sha(dst)
        entries.append({'id': vid, 'title': label, 'path': '视频附件/' + dst.name, 'relative_path': '视频附件/' + dst.name,
                        'source_relative_path': src.relative_to(ROOT).as_posix(), 'sha256': sha(dst), 'bytes': dst.stat().st_size})
    write_json(OUT / 'media_manifest.json', {'entries': entries, 'embedded_in_docx': True,
               'scope': 'Video copies and original files have identical bytes; playback is by double-clicking embedded Package attachments.'})
    return entries


def plots():
    fig, ax = plt.subplots(1, 2, figsize=(9.1, 3.1), constrained_layout=True)
    steps = [1000, 2000, 3000, 4000, 5000]
    ax[0].plot(steps, [.04952885,.04072450,.03388455,.03160163,.03000003], 'o-', color='#226B92')
    ax[0].set(title='训练损失持续降低', xlabel='新增优化器更新次数', ylabel='最近100步训练loss均值')
    ax[1].plot(steps, [1.965089,2.137033,.953194,1.307184,1.169662], 'o-', color='#93602D')
    ax[1].scatter([3000], [.953194], s=90, facecolors='none', edgecolors='#27733D', linewidths=2)
    ax[1].set(title='验证误差没有单调降低', xlabel='新增优化器更新次数', ylabel='验证动作MAE  度')
    for a in ax:
        a.set_xticks(steps, ['1k','2k','3k','4k','5k']); a.grid(axis='y', alpha=.2); a.spines[['top','right']].set_visible(False)
    fig.savefig(WORK / 'training.png', dpi=170); plt.close(fig)
    fig, a = plt.subplots(figsize=(7.1, 4.2), constrained_layout=True)
    grid = [(x,y) for x in range(-40,41,20) for y in range(-40,41,20)]
    a.scatter(*zip(*grid), marker='s', s=32, facecolors='none', edgecolors='#6F777B', label='25个合成训练位置')
    a.scatter([-10,-10,10,10],[-10,10,-10,10], marker='x', s=50, color='#CF8E35', label='4个验证位置')
    pts = [(int(round(r['offset_xy_m'][0]*1000)),int(round(r['offset_xy_m'][1]*1000))) for r in H['latest_final8']]
    a.scatter(*zip(*pts), s=43, color='#23885F', label='8个最终测试位置')
    for i,(x,y) in enumerate(pts):
        a.annotate(f'{i:02d}', (x,y), xytext=(4,4), textcoords='offset points', fontsize=8)
    a.scatter([10],[0],marker='*',s=140,color='#83539B',label='用户追加位置')
    a.set(xlim=(-48,48), ylim=(-48,48), xlabel='相对参考位置X偏移  mm', ylabel='Y偏移  mm')
    a.set_aspect('equal'); a.grid(alpha=.16); a.legend(loc='upper left',bbox_to_anchor=(1.01,1),fontsize=7.8)
    fig.savefig(WORK/'positions.png',dpi=180); plt.close(fig)


def main():
    global doc, media
    OUT.mkdir(parents=True, exist_ok=True); WORK.mkdir(parents=True, exist_ok=True)
    media = build_media(); plots()
    doc = Document()
    for border in doc.styles.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
    sec = doc.sections[0]
    sec.page_width = Cm(21); sec.page_height = Cm(29.7)
    sec.top_margin = Cm(1.8); sec.bottom_margin = Cm(1.7)
    sec.left_margin = sec.right_margin = Cm(2)
    sec.header_distance = sec.footer_distance = Cm(.8)
    for name,size in [('Normal',10.5),('Title',25),('Heading 1',16),('Heading 2',12),('Heading 3',10.5)]:
        s=doc.styles[name]; s.font.name=FONT; s.font.size=Pt(size); s.font.color.rgb=RGBColor(0,0,0)
        s._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),FONT)
        s.paragraph_format.space_after=Pt(7)
        if name.startswith('Heading'):
            s.font.bold=True; s.paragraph_format.space_before=Pt(8)
    hp=sec.header.paragraphs[0]; set_font(hp.add_run('ACT 项目记录  /  第一阶段仿真归档'),8,color='666666')
    fp=sec.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    set_font(fp.add_run('2026-10-24  ·  第 '),8,color='666666')
    fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); fp._p.append(fld)
    set_font(fp.add_run(' 页'),8,color='666666')
    doc.core_properties.title='ACT 抓放仿真第一阶段复盘与真机交接'
    doc.core_properties.subject='固定场景下的局部位置泛化、可追溯实验与下一阶段交接'
    doc.core_properties.author='ACT项目组'

    doc.add_heading('ACT 抓放仿真\n第一阶段复盘与真机交接',0)
    para('记录范围  2026年10月23日至24日  |  台式机本地仿真',11,False,12)
    heading('本阶段得到的结果')
    para('我们已把“机械臂看起来在动”的预览，推进到真正由ACT根据双路图像和关节状态产生动作、经物理接触把红块放入灰盒的闭环。在固定模型、相机与道具条件下，最终策略通过4个验证位置、8个事先保留的测试位置；用户随后追加的另1个位置也单次成功。')
    table(['验收对象','保存结果','可支持的结论'],[
        ('原位与验证','原位1/1；验证4/4','选模型前的任务准入'),
        ('最终8个位置','8/8；抬升96.78–105.94 mm','本场景内未用于选模的XY位置测试'),
        ('用户追加1点','(+10,0) mm；1/1','追加插值位置验证，单独统计'),
        ('动作来源核验','8点与追加点全部重新forward，最大差0','保存动作确由冻结ACT和所存观察产生'),
        ('真机执行','本轮硬件命令0','尚未证明真机抓放成功'),
    ],[3.1,5.8,8.1])
    heading('最终保留的策略')
    para('采用新5000步训练分支中的step3000检查点，配合ensemble1逐帧查询与时间集成。5000步是训练预算，不是“最后一步必然最好”；最终选择先看原位和验证位置的实际抓放，再看稳态抓持，离线MAE仅用于同分排序。')
    para('权重身份  c25ffa08…c506c2c。来源为公开权重→真实数据续训1000步→sim300→新分支3000步。沿这条分支仿真适配累计3300次，真实数据续训1000次另记，不累加其他实验分支。',9)
    para('结果边界  这是固定红块、固定盒子、固定相机、固定物理参数下，±40 mm局部区域内的XY插值能力；尚未检验真实D405画面、实机动力学、旋转/遮挡/光照变化。',10.5,True)
    para('本报告内含14段MP4附件。双击视频图标可调用播放器；若Office限制附件启动，打开随报告提供的“视频附件”文件夹。附件是完整视频本体，并非只放网页链接。',9)
    sources('history_evidence.json；final_selection；final8；extra_summary；全部附件SHA见media_manifest.json。')

    page('1  起点与完整执行框架')
    para('硬件目标是Seeed reBot B601-RS，腕部D405与顶部D435i。此前已在笔记本完成上位机通信、保持、回零与支撑确认后失能；本阶段在台式机建立MuJoCo仿真，避免把尚未验证的策略直接交给重型机械臂。')
    table(['来源','本项目实际采用的内容'],[
        ('社区公开资源','nikodembartnik/rebot_cube数据：50 episodes、20,948帧、30 Hz；对应公开act_rebot_cube权重。它不是Seeed官方发布的ACT权重。'),
        ('Seeed开源资源','固定commit 438f7713…3997a的reBot_RS_Sim模型资产，共44项。用于机械结构和仿真重建。'),
        ('本机仿真环境','MuJoCo 3.17.0；ACT环境LeRobot 0.6.1、Torch 2.10.0+cu130、Python 3.12。仿真和Torch环境分开。'),
        ('本机资源约束','i5-13600KF、32 GB内存、RTX 5060 Ti 8 GB；进程限定4个逻辑核、低优先级、Torch 2线程。'),
    ],[3.5,13.5])
    heading('ACT如何从观察变成抓放')
    table(['步骤','输入或输出','职责'],[
        ('① 采样','当前7维状态＋顶部RGB＋腕部RGB','每张图3×480×640；没有物体真值坐标或教师动作'),
        ('② 处理器','沿用保存的归一化配置与统计','保持公开模型、仿真续训、推理的数据语义一致'),
        ('③ ACT预测','一次输出未来100×7动作','用图像和状态估计动作块，不调用IK去替学生抓取'),
        ('④ 时间集成','ensemble1汇总覆盖当前时刻的预测','每个动作帧重新观察、重新forward、加权得到7维目标'),
        ('⑤ 仿真执行','目标经仿真坐标转换→PD力矩→物理步进','机器人和方块靠接触运动，物体不绑到夹爪上'),
        ('⑥ 独立判定','qpos/qvel、双指接触、盒内位置与速度','将“发出了动作”和“确实抓起并放稳”分开'),
    ],[2.4,6.2,8.4],8.8)
    para('这里修改的是数据覆盖、训练适配、执行循环和验证方式；没有另造一套ACT网络。迁往真机时，①至④的语义需要保留，⑤必须接真实控制器，不能把仿真PD或仿真关节映射直接当作电机驱动。')
    sources('download；real_data_audit；upstream；train5000；worker/query/controller代码锚点。')

    page('2  先把仿真任务做成立')
    heading('问题一  网页能播放不等于有抓取物理',2)
    para('最初先导入机器人模型、关节映射和相机，回放403帧状态并导出100×7单次预测。该阶段21项检查通过，但记录明确标注：未验证物理控制、ACT闭环和抓取成功。直接修改qpos的预览只能看姿态，不能证明力、摩擦、碰撞或抓持成立。')
    para('处理方式：补上力矩伺服、重力和碰撞接触，让机械臂经MuJoCo物理步进运动；评分重新读取实际物体状态。之后才有资格比较ACT策略。')
    heading('问题二  公开策略在重建场景中没有直接成功',2)
    para('改成红块入灰盒后，公开权重及本机真实数据续训1000步的权重，各运行15秒、真实查询5次，均未抓起；参考示范回放最大抬升仅0.258 mm。真实数据画面、坐标、夹爪语义与近似仿真之间的差异，并不会因机械臂型号相同就自动消失。')
    para('处理方式：先构造可物理完成任务的合成教师，得到403帧、30 Hz的成功示范。平顶40 mm版本教师抬升99.040 mm并入盒，用于后续仿真适配。教师有明确轨迹目标；教师成功仅证明任务可完成，不能写成ACT成功。')
    heading('问题三  道具反复改变会掩盖算法问题',2)
    para('按用户约定去掉红块顶部凸点，冻结道具尺寸，另行导出STL与STEP。后续围绕固定几何优化策略，不用反复更换尺寸挑选成功画面。')
    table(['道具','冻结尺寸','含义'],[
        ('红色方块','40×40×20 mm','实心、平顶、无凸点，俯视正方形'),
        ('灰色开口盒','外120×120×50 mm','底厚4 mm、壁厚4 mm'),
        ('盒内腔','112×112×46 mm','模型尺寸推导值，供装入判定使用'),
    ],[3.1,5.5,8.4])
    para('以上是本项目的道具约定，并非数据作者提供的实物测量标准。相机外观拟合也不等于已完成真实D405手眼标定。',9)
    sources('first_pass；task_failures；teacher_flat；props/props_verification。')

    page('3  单点成功之后为何换位置失败')
    heading('先解决动作块被反复截断的问题',2)
    para('在物理成功示范上新增300次更新，batch=2、lr=1e-5、FP32，约92.93秒。最初每次只执行10步就重新预测，会不断抛弃动作块后面的闭爪动作，机械臂在抓取前的相位停滞。读取保存预测后，同一模型、同一场景恢复原定100步执行长度，固定点即完成抓放：抬升94.691 mm，15秒内真实ACT查询5次。')
    para('这一轮说明动作块长度与执行相位必须匹配；不能推导为“所有任务都必须一次执行100步”。长期不看新图像又会让位置纠错变慢，后续需要另一种执行方式。')
    video('V01','单位置适配后的ACT成功抓放')
    heading('用户提出的换位置测试暴露了局部拟合',2)
    para('保持模型、相机、机械臂初态和场景不变，仅修改方块初始XY，sim300原位1/1，5个新位置0/5，最大抬升仅2.087 mm且没有双侧共同夹持。把公开模型、desktop1000和sim300放在同一循环比较的18试次，也只有sim300原位成功。')
    para('初始关节状态相同、图像改变时，ACT预测确有变化。因此不能把失败直接说成“完全没看图/没识别红色”；更准确的判断是：它对图像有响应，但单点示范及重建场景下的动作覆盖不足以产生有效抓持。灾难性遗忘是否是主因，现有证据不能单独确定。')
    video('V02','原位置与五个新位置同步回看')
    heading('这一步留下的经验',2)
    para('固定点成功后立即冻结模型再换位置，比继续观察同一条成功轨迹更能发现问题。要保留失败，不把教师动画、网络前向输出和物理成功混为一谈。')
    sources('sim300_train；prefix_failure；sim300_success；generalization；checkpoint_comparison。')

    page('4  扩大数据覆盖也曾经没有奏效')
    para('针对单点覆盖不足，从同一条403帧成功轨迹出发，在±40 mm网格内用六关节Jacobian IK生成25个XY位置变体；4个位置单独验证。训练共10,075帧，验证1,612帧。它们是合成变体，不是25次独立人工示范。')
    para('40条合成教师抓放均成功，附加穿透检查严格通过39/40；未严格通过的是旧诊断点，早期穿透约1.020 mm，并非剔除了一条训练示范。后续增加22条恢复专家，合并成47个训练回合、14,695帧；这些成功仍只属于教师。')
    table(['10月23日分支','新增更新与来源','保存结果'],[
        ('v1','从sim300新增600；batch2，lr3e-6','原位准入，验证2/4，作为当时保留模型'),
        ('v2','从v1新增1200；step600是中间点','600验证2/4但原位失败；1200验证0/4，MAE约1.535°'),
        ('clean v3','从v1另开2400步分支','选step1800，验证0/4，未保留'),
        ('recovery v1','从v1另开2400；lr1e-6','step1800和2400均验证0/4、原位失败'),
    ],[3.0,6.3,7.7],8.8)
    para('旧v1 step600随后六个保留位置0/6、五个诊断位置1/5。低离线动作误差并未转成闭环成功，扩大数据或继续更新也没有自动解决抓偏后的反馈问题。分支之间共享来源，但步数不能累加成“一次连续长训”。')
    video('V03','旧多位置保留模型的六点测试结果')
    heading('由此改变的处理思路',2)
    para('不再只盯训练loss。明确设置连续训练预算，保存多个检查点，在同一物理循环中比较“权重＋执行模式”，把原位、验证位置和最终测试位置分开。恢复数据这条路线保留源码与失败记录，但没有把它包装成最终成功的原因。')
    para('恢复轨迹聚合只借鉴了DAgger思路；本项目未进行在线逐步专家标注，不能声称完整复现DAgger。',9)
    sources('teacher_positions；teacher_positions_physics；teacher_trace_audit；adapt_v1/v2/v3；adapt_recovery_v1；old_final6。')

    page('5  连续5000步与执行循环一起迭代')
    para('10月24日从sim300另开新分支，固定原25/4数据划分，一次连续更新5000步，每1000步存检查点。此前旧step600没有接入这条新分支。')
    table(['配置','实际值'],[
        ('优化器与精度','AdamW；FP32；lr 1e-5；weight decay 1e-4；梯度范数裁剪10'),
        ('batch与样本量','全程batch8；OOM0；40,000次采样观察窗口，非40,000条独立示范'),
        ('归一化与数据隔离','保留来源处理器与统计；训练不读取测试轨迹、图像和动作标签'),
        ('时间与显存','循环含验证/保存1201.80秒；进程约1221.85秒；峰值allocated3766.48 MiB，reserved4118 MiB'),
        ('资源控制','4个逻辑核、低优先级、Torch2线程；GPU监测约55–68°C；不是温度硬保护'),
    ],[3.7,13.3],8.9)
    picture(WORK/'training.png',16.5,'图1  同一连续分支的保存点。训练loss下降，不代表验证误差或物理抓放单调改善。')
    table(['新增步数','训练loss 最近100步均值','归一化验证MAE','验证MAE 度'],[
        ('1000','.049529','.070125','1.965'),('2000','.040725','.079688','2.137'),
        ('3000','.033885','.035402','.953'),('4000','.031602','.048672','1.307'),('5000','.030000','.044748','1.170'),
    ],[2.5,6,4.5,4],8.7)
    para('执行循环新增独立CUDA worker。chunk100每15秒仿真仅5次真实forward；ensemble1每个30 Hz动作帧都重新取双RGB和state，15秒真实forward450次，再对覆盖当前时刻的预测加权。改变的是观察和执行的组织方式，网络本身仍是ACT。',10)
    sources('train_preflight；train5000；cuda_preflight；long_train/query/ensemble代码锚点。')

    page('6  如何选模型并处理越界中止')
    para('按预定顺序选择：同模式原位成功才准入，再比较4个验证位置的完整物理成功数、稳态抓持数，离线MAE只用于同分。没有依据最终8点成绩重新选权重。')
    table(['候选权重','chunk100 原位／验证','ensemble1 原位／验证'],[
        ('旧保留step600','通过／2/4','通过／1/4'),('新1000','通过／2/4','失败／0/4'),
        ('新2000','越界中止／未运行','通过／4/4'),('新3000','通过／验证第4点越界','通过／4/4，最终选定'),
        ('新4000','越界中止／未运行','通过／3/4'),('新5000','越界中止／未运行','越界中止／未运行'),
    ],[3.0,7,7],9)
    para('step2000和step3000的ensemble1均达到4/4；step3000验证归一化MAE .0354018更低，按同分规则胜出。step3000的chunk100前3个完整结果依次为失败、成功、成功，第4个在360/450帧越界，不能当作完整验证通过。')
    heading('错误必须分清楚，不能靠放宽保护制造成功',2)
    para('早期矩阵遇到原始夹爪目标低于0而停下。后来保留原runner和原中止证据，新增v2 runner，保存raw chunk、原生关节目标和首次未施加的目标，并证明此前施加值均合法。仅能复现为原validate_target同一种异常的情况，归为policy_limit_abort；没有clip，没有放宽边界，未知错误仍停。')
    table(['统计类别','实际次数','如何记账'],[
        ('完整物理回合','39','按完整评分判断成功或失败'),('越界后的部分回合','5','保留部分轨迹；不算完整成功'),
        ('对应原位越界后未尝试','16','记录未运行/null，不伪写成0/4'),('原计划合计','60','39＋5＋16；真实ACT查询11,665次'),
    ],[5.4,2.2,9.4])
    para('5次越界均为第7通道夹爪目标负值：2000/chunk原位352帧，3000/chunk验证360帧，4000/chunk原位352帧，5000/chunk原位355帧，5000/ensemble原位328帧。即便首次只有约−0.0339 mm，也按原硬边界拒绝；这不等于验证了真机夹爪限位。',9)
    sources('matrix_protocol；matrix_original_stop；matrix_amendment；matrix_completion；proposal；final_selection。')

    page('7  最终位置验证得到什么')
    para('先冻结模型、模式和hash，再运行8个保留位置，每点一次。它们不与本地25训练、4验证、5诊断和6旧测试点重合，均在±40 mm区内。留出只针对本地实验；原50条真实示范是否含相近位置未知。')
    picture(WORK/'positions.png',13.5,'图2  位置划分。坐标均相对同一参考点；紫色为看过8点结果后由用户要求追加的单点。')
    rows=[]
    for r in H['latest_final8']:
        x,y=[int(round(v*1000)) for v in r['offset_xy_m']]
        rows.append((r['id'].replace('newtest_',''),f'({x:+d}, {y:+d})','成功',f"{r['max_lift_mm']:.2f}",r['longest_elevated_bilateral_frames']))
    table(['编号','XY偏移 mm','完整放入','最大抬升 mm','连续抬高双指接触 帧'],rows,[1.5,4.1,3,3.4,5],8.6)
    para('成功判定同时要求：抬升至少30 mm；方块8个角点的XY投影全在盒内；落在盒底；终速小于0.05 m/s；末态无手指接触。独立审计还要求抬高至少30 mm且双指接触连续至少3帧。8点最长连续抬高双指接触68–73帧，均通过。',10)
    para('追加(+10,0) mm也成功1/1：抬升101.374 mm、双指共同接触97帧、其中连续抬高70帧。它距原8点最近25.495 mm，未与48个已见唯一偏移重合；是在旧区域内另选点，不是范围外泛化。',10)
    sources('new_plan；final_selection；final8；extra_registration；extra_summary；novelty。')

    page('8  如何排除直接播放轨迹的嫌疑')
    heading('保留每一步真实输入和输出',2)
    para('最终8点共保存3600份双路观察、当前状态、原始100×7预测块、集成权重来源、施加目标、qpos/qvel及接触。策略worker只读图像与状态，没有读取测试标签、物体真实坐标、教师动作或IK目标。')
    table(['核验','保存证据与结果'],[
        ('重新调用同一网络','8点3600次fresh forward，2,520,000个动作值逐值相同，最大差0；追加点450次、315,000值也最大差0。'),
        ('动作施加链','独立重算时间集成及坐标转换，与全部实际施加目标逐值一致；查询state等于当时实际动作前状态。'),
        ('视觉变化响应','8点初始关节状态完全一致；28对初始动作块均有差异，最大绝对差范围1.2702–7.6057°；至少顶部图像变化影响了动作。'),
        ('物理抓取真实性','从qpos/qvel与双指接触重算评分一致；没有将方块绑在手上，没有中途瞬移到盒内。'),
    ],[3.6,13.4],9)
    para('这些证据足以说明保存动作由当前观察下的ACT产生，并非直接执行一条固定动作文件。但“输入改变导致输出改变”还不等于证明模型理解了颜色、形成了通用目标检测或具备任意场景泛化。')
    picture(ROOT/'runs/mujoco_validation_20261023/additional_position_20261024/actual_inputs_frame221.png',11,'图3  追加位置的实际输入与场景回看。完整清晰画面见450帧视频。')
    video('V13','追加位置实际输入与完整抓放视频')
    sources('final8_fresh；initial_visual；final_completion；extra_fresh；extra_summary。')

    page('9  实时性与独立基线的边界')
    heading('30 Hz仿真时间不等于真机30 Hz',2)
    para('本机CUDA预检使用旧step600与固定观察，20次暖启动worker内部预测路径平均19.153 ms，包含处理器、网络和拷回CPU；加上压缩/文件/IPC的往返平均48.200 ms，尚不含渲染和物理步进。这不是最终8点的吞吐测量。同步推理会暂停仿真时间，不能据15秒视频宣称真实墙钟端到端30 Hz。')
    para('ensemble1沿用原版按旧预测在前的exp(−0.01×序号)加权。稳定覆盖100份预测时，最新观察预测仅占约0.585%，平均预测来源年龄57.70动作帧、约1.923秒。每帧重新forward并不等于只执行最新预测第一步；平滑与纠偏速度有取舍。')
    table(['测量对象','实测','能说明什么'],[
        ('旧step600 CUDA内部预测','平均约19.15 ms','包含处理器、网络、拷回CPU'),
        ('同一预检CUDA往返','平均约48.20 ms','包含当前IPC，不含渲染/物理'),
        ('新交接包CPU前向流程','3次平均约395.7 ms','可离线复算；不足以证明实时控制'),
    ],[5.4,4.2,7.4])
    heading('ALOHA用于检查另一套配套链路',2)
    para('另外使用LeRobot公开ALOHA transfer cube human ACT和gym-aloha配套环境，CPU执行5个seed回合，4成功、1失败，真实网络forward共17次。它是14维状态、单顶部相机的双臂递交任务；动作50 Hz但每100步才查询一次，不能把50 Hz说成网络每秒推理50次。')
    para('这个5-seed短测用于确认模型与环境配套时的ACT链路；它既不等于公开500回合83%指标的复现，也不属于B601抓放成绩。环境独立使用MuJoCo 3.2.7，不能与B601的3.17.0环境混装。')
    video('V14','ALOHA基线seed20261024的一回合视频')
    para('本阶段同时改变训练预算、batch、学习率、CUDA执行与反馈模式，没有完整析因实验。成绩提升可以归因于这组改进的联合结果，不能声称仅由“多训2000步”或“每帧看图”单独造成。',10.5,True)
    sources('cuda_preflight；ensemble；aloha_protocol；aloha；aloha_audit；stage2 offline报告。')

    page('10  经验沉淀与源码落点')
    table(['曾遇到的问题','本轮有效处理思路','下次沿用的做法'],[
        ('动画像成功，物理未验证','加入力矩伺服、接触和独立评分','先做任务成立验证，再谈模型成绩'),
        ('任务道具和镜头反复改变','冻结平顶40 mm道具与相机条件','改变条件时另开实验，不覆盖旧证据'),
        ('固定点好看，挪位置失败','扩大合成位置覆盖并留出测试','先冻结模型，再打开测试位置'),
        ('动作块后半程一直没执行','查预测中的闭爪相位并修执行长度','同时查预测、执行模式与物理轨迹'),
        ('loss下降但抓取变差','比较多个checkpoint的物理闭环','先原位/验证任务成功，再以MAE同分排序'),
        ('夹爪越界使整轮中断','证明并分类已知limit异常','保留原值和首次拒绝，不clip伪造成功'),
        ('担心固定轨迹或隐藏专家','保存全部观察并fresh网络复算','动作身份、施加链和物理评分各自核验'),
    ],[4.3,6.1,6.6],8.6)
    heading('关键源码入口与冻结行号',2)
    selected=[('scene','冻结道具与机器人场景'),('controller','仅仿真PD/力矩伺服'),('teacher','合成单示范导出'),('grid','位置划分'),('ik','合成位置轨迹变体'),('long_train','连续5000步适配'),('query','新鲜观察下的闭环评估'),('ensemble','时间集成'),('fresh_verify','保存观察重算预测'),('score','物理成功判定')]
    table(['功能','源码位置'],[(label,f"{H['code_anchors_current_at_archive'][key]['relative_path']}:{H['code_anchors_current_at_archive'][key]['start_line']}") for key,label in selected],[4.6,12.4],8.3)
    p=para('GitHub源码里程碑  '+COMMIT[:12],9)
    hyperlink(p,'  查看固定提交','https://github.com/1595530254a-afk/act/commit/'+COMMIT)
    para('公开复现入口：docs/stage1_20261024/public_reproduction.md。它包含最小冻结输入的恢复位置、路径/hash重算和环境分离步骤。Git保存集成源码与精选验证视频；大权重由笔记本包携带，官方模型资产按固定清单下载并保留来源许可。',9)
    sources('code_anchors_current_at_archive；source_manifest.json；git_scope_audit；public_reproduction.md。')

    page('11  第二阶段如何交给笔记本执行')
    para('这一阶段交付的是“更好的仿真策略候选＋同语义的推理代码＋明确操作单”。真机仍使用已验证的驱动和保持→回零→支撑确认→失能链。仿真成功不能替代真实关节单位、方向、零位和相机输入的一致性。')
    table(['顺序','笔记本Codex执行','完成后应留下什么'],[
        ('1 离线复算','并行解压新包，用现有ACT Python运行纯测试与默认离线脚本；再试CUDA。','7项测试结果、3次新forward、模型hash、实际耗时；硬件命令0。'),
        ('2 保留已有安全链','沿用已实际成功的保持/回零/承重确认/失能程序；包不替换驱动。','原控制器路径及当前映射版本；硬件未变时不重问已确认事实。'),
        ('3 读取真实观察','只读枚举D405和D435i并固定序列号；由已有控制器导出真实7通道反馈。','双RGB、反馈单位与顺序、映射ID、本机单调时钟接收时间。'),
        ('4 只观察shadow','常驻ACT收到ready后再提供新快照，保存原始chunk和集成动作；不向电机发送。','至少30秒的真实间隔、输入年龄、端到端耗时、动作范围；不以completed冒充实时验收。'),
        ('5 回传再接动作','按包内RETURN_stage2.json与返回清单打包；依据真实读数接已有动作适配器。','只补差异，不重做整个项目复盘；首次真机策略动作另行限域验证。'),
    ],[2.2,8.4,6.4],8.7)
    heading('包里携带什么',2)
    para('选定step3000的六个完整模型/处理器文件；portable NumPy时间集成；3个冻结观察/参考预测；默认离线推理；RealSense只读枚举和单次快照；真实状态数据合同；给Codex的中文操作单和返回模板。没有CAN写入口，没有机械臂SDK，没有仿真joint_mapping或PD控制。')
    para('离线复算命令（把ACT_PYTHON换成笔记本已有解释器）',9,True)
    para('ACT_PYTHON runtime/test_stage2.py\nACT_PYTHON runtime/stage2_policy.py --device cpu',9)
    para('已验证：新包7项纯测试通过；3次CPU真模型forward与保存CUDA参考最大差0.00004578°；时间集成与冻结仿真450帧逐值一致。CPU每次约396 ms，只是可运行证明。窗口建议同时保留推理日志、顶部/腕部画面和现有硬件控制台，便于人工同步核对。')
    para('D405标定说明：纯图像ACT不会自动要求一个显式手眼矩阵，但真实镜头、裁剪、曝光和安装位置必须与策略语义相容；若后续用深度转机器人坐标或几何规划，再单独完成相应标定。不要拿仿真拟合参数当成实机标定。',9)
    sources('deployment/stage2_20261024/先读我_给笔记本Codex.md；NOTICE.md；snapshot_contract.json；runtime/test_stage2.py。')

    page('附录一  最终验证视频 00至03')
    para('下列视频是报告内部的MP4附件；图标双击打开。V04为八点同步总览，V05至V12为各自实际双路输入与场景回看，均为450帧、30 fps、15秒。视频播放时长不代表运行耗时。',10)
    video('V04','最终8点同步总览')
    for i in range(4):
        r=H['latest_final8'][i]; x,y=[int(round(v*1000)) for v in r['offset_xy_m']]
        video(f'V{i+5:02d}',f"位置{i:02d}  ({x:+d}, {y:+d}) mm  成功，最大抬升{r['max_lift_mm']:.2f} mm")
    para('完整文件名、原视频相对路径、大小与SHA256见media_manifest.json。视频附件为原文件逐字节副本；原视频继续保存在 D:\\外置大脑\\ACT\\视频。',9)
    sources('final8 evaluation SHA256 6ef257e217f20b51412dd82fbb80a590103cf97b08c47694b6656dc271ba8c3b。')

    page('附录二  最终验证视频 04至07与归档索引')
    for i in range(4,8):
        r=H['latest_final8'][i]; x,y=[int(round(v*1000)) for v in r['offset_xy_m']]
        video(f'V{i+5:02d}',f"位置{i:02d}  ({x:+d}, {y:+d}) mm  成功，最大抬升{r['max_lift_mm']:.2f} mm")
    heading('报告包内的长期记录',2)
    table(['文件','用途'],[
        ('本Word文档','完整阶段记录与14个视频内嵌附件'),
        ('视频附件/V01.mp4至V14.mp4','附件不能启动时的独立播放器备份'),
        ('media_manifest.json','原始视频来源、逐文件大小和SHA256'),
        ('evidence/history_evidence.json 与 .md','逐问题的原始工件、源码函数行号和82文件hash'),
        ('evidence/source_manifest.json','公开源码、模型身份、环境及精选视频清单'),
        ('evidence/public_reproduction.md','Git源码恢复与最小仿真输入说明'),
    ],[7.0,10.0],8.6)
    para('选定权重 SHA256',9,True)
    para('c25ffa08fd2b67343f37eab06bb948642e2a6d2578e387a563ba81105c506c2c',8.5)
    para('第一阶段保留原始失败与成功证据，本次整理未增加训练或物理回合。第二阶段从真实观察接入和只观察推理开始，候选策略经过真实语义与时延验证后，再进入受控真机抓放。',10)

    base = WORK/'stage1_base.docx'; doc.save(base)
    ev=OUT/'evidence'; ev.mkdir(exist_ok=True)
    for src,name in [(EVIDENCE,'history_evidence.json'),(EVIDENCE.with_suffix('.md'),'history_evidence.md'),
                     (ROOT/'docs/stage1_20261024/source_manifest.json','source_manifest.json'),
                     (ROOT/'docs/stage1_20261024/public_reproduction.md','public_reproduction.md')]:
        shutil.copy2(src,ev/name)
    write_json(WORK/'build_record.json',{'base':str(base),'source_commit':COMMIT,'media_count':len(media),
               'media_total_bytes':sum(v['bytes'] for v in media),'history_evidence_sha256':sha(EVIDENCE),
               'CPU_BUDGET':CPU_BUDGET,'training_steps':0,'physics_steps':0,'hardware_commands':0})
    print(json.dumps({'base':str(base),'output_directory':str(OUT),'media_count':len(media),'source_commit':COMMIT},ensure_ascii=False))


if __name__ == '__main__':
    main()
