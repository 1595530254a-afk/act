"""Generate the simulation teaching report without video links or attachments.

Author: 请输入名字QaQ (1595530254a-afk).
Uses only the checked-in teaching text and compact plot data, never model IO.
"""
import argparse
import hashlib
import json
from pathlib import Path
import re
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'scripts'))
from cpu_budget import CPU_BUDGET

import matplotlib
matplotlib.use('Agg')
from matplotlib import pyplot as plt
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT = 'Microsoft YaHei'
AUTHOR = '请输入名字QaQ'
plt.rcParams.update({'font.sans-serif': [FONT], 'axes.unicode_minus': False, 'font.size': 10})


def font(run, size=11, bold=False, name=FONT):
    run.font.name = name; run.font.size = Pt(size); run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    f = run._element.get_or_add_rPr().rFonts
    f.set(qn('w:eastAsia'), FONT); f.set(qn('w:ascii'), name); f.set(qn('w:hAnsi'), name)


def runs(paragraph, text, size=11):
    pattern = r'(\[[^\]]+\]\(https?://[^)]+\)|\*\*[^*]+\*\*)'
    for piece in re.split(pattern, text):
        if not piece:
            continue
        match = re.fullmatch(r'\[([^\]]+)\]\((https?://[^)]+)\)', piece)
        if match:
            label, target = match.groups()
            rel = paragraph.part.relate_to(target, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
            node = OxmlElement('w:hyperlink'); node.set(qn('r:id'), rel)
            r = OxmlElement('w:r'); p = OxmlElement('w:rPr')
            for tag, value in [('color','245C88'), ('u','single'), ('sz',str(int(size*2)))]:
                el = OxmlElement('w:' + tag); el.set(qn('w:val'),value); p.append(el)
            fs = OxmlElement('w:rFonts'); fs.set(qn('w:eastAsia'),FONT); p.append(fs)
            r.append(p); t = OxmlElement('w:t'); t.text = label; r.append(t); node.append(r); paragraph._p.append(node)
        elif piece.startswith('**'):
            font(paragraph.add_run(piece[2:-2]), size, True)
        else:
            font(paragraph.add_run(piece), size)


def plots(data, output):
    steps = [r['step'] for r in data['checkpoints']]
    fig, axes = plt.subplots(1, 2, figsize=(9,3.05), constrained_layout=True)
    for ax, values, title, ylabel, color in [
        (axes[0],[r['recent100_loss_mean'] for r in data['checkpoints']],'训练loss持续下降','最近100步训练loss均值','#246D94'),
        (axes[1],[r['validation_degree_mae'] for r in data['checkpoints']],'验证动作误差并非单调改善','MAE  度','#99652F')]:
        ax.plot(steps, values, 'o-', color=color); ax.set(title=title,xlabel='新增优化器更新次数',ylabel=ylabel)
        ax.set_xticks(steps,['1k','2k','3k','4k','5k']); ax.grid(axis='y',alpha=.2); ax.spines[['top','right']].set_visible(False)
    axes[1].scatter([3000],[data['checkpoints'][2]['validation_degree_mae']],s=90,facecolors='none',edgecolors='#2B824E',linewidths=2)
    fig.text(.995,.005,AUTHOR+'  ·  2026',ha='right',fontsize=7)
    fig.savefig(output/'training.png',dpi=180); plt.close(fig)
    fig,ax = plt.subplots(figsize=(7.2,3.65),constrained_layout=True)
    grid=[(x,y) for x in (-40,-20,0,20,40) for y in (-40,-20,0,20,40)]
    ax.scatter(*zip(*grid),marker='s',s=32,facecolors='none',edgecolors='#758087',label='25个合成训练位置')
    ax.scatter([-10,-10,10,10],[-10,10,-10,10],marker='x',s=50,color='#BD812A',label='4个验证位置')
    points=[tuple(v*1000 for v in r['offset_xy_m']) for r in data['final8']]
    ax.scatter(*zip(*points),s=38,color='#23805C',label='8个最终测试位置')
    for i,(x,y) in enumerate(points):ax.annotate(f'{i:02d}',(x,y),xytext=(4,4),textcoords='offset points',fontsize=8)
    ax.scatter([10],[0],marker='*',s=130,color='#83539A',label='用户追加1个位置')
    ax.set(xlim=(-48,48),ylim=(-48,48),xlabel='相对参考点X偏移  mm',ylabel='Y偏移  mm')
    ax.set_aspect('equal'); ax.grid(alpha=.15); ax.legend(loc='upper left',bbox_to_anchor=(1.01,1),fontsize=8)
    fig.text(.995,.005,AUTHOR+'  ·  2026',ha='right',fontsize=7)
    fig.savefig(output/'positions.png',dpi=180); plt.close(fig)


def make_table(doc, lines):
    cells = [[v.strip() for v in s.strip().strip('|').split('|')] for s in lines]
    rows=[cells[0]]+cells[2:]
    count=len(rows[0]); widths={2:[4.0,13.0],3:[3.2,5.1,8.7],4:[2.8,5.5,4.3,4.4]}[count]
    if rows[0][0]=='数据划分':widths=[2.5,9.5,5.0]
    table=doc.add_table(rows=0,cols=count); table.autofit=False
    for i,w in enumerate(widths):table.columns[i].width=Cm(w)
    for n,values in enumerate(rows):
        row=table.add_row(); row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        if n==0:row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        for i,value in enumerate(values):
            cell=row.cells[i]; cell.width=Cm(widths[i]); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=cell._tc.get_or_add_tcPr(); margin=OxmlElement('w:tcMar')
            for side in ['top','bottom','left','right']:
                el=OxmlElement('w:'+side); el.set(qn('w:w'),'70');el.set(qn('w:type'),'dxa');margin.append(el)
            pr.append(margin)
            if n==0:
                fill=OxmlElement('w:shd');fill.set(qn('w:fill'),'E8EEF2');pr.append(fill)
            p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(1.5);p.paragraph_format.line_spacing=1.1
            if n==0:font(p.add_run(value),9.5,True)
            else:runs(p,value,9.5)
            if i==0 and count==4:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        b=OxmlElement('w:'+side);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');borders.append(b)
    table._tbl.tblPr.append(borders)
    end=doc.add_paragraph();end.paragraph_format.space_after=Pt(1);end.paragraph_format.line_spacing=.5


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output-dir',type=Path,default=ROOT/'output/ACT仿真教学复盘_20261024')
    args=parser.parse_args(); output=args.output_dir.resolve();output.mkdir(parents=True,exist_ok=True)
    source=ROOT/'docs/stage1_20261024/ACT仿真训练复现与调参复盘.md'
    dataset=ROOT/'docs/stage1_20261024/教学图表数据.json'
    data=json.loads(dataset.read_text('utf-8'));plots(data,output)
    doc=Document()
    for border in doc.styles.element.xpath('.//w:pBdr'):border.getparent().remove(border)
    sec=doc.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    sec.top_margin=Cm(1.7);sec.bottom_margin=Cm(1.65);sec.left_margin=sec.right_margin=Cm(2)
    sec.header_distance=sec.footer_distance=Cm(.7)
    for name,size in [('Normal',11),('Title',23),('Heading 1',15.5),('Heading 2',12),('Heading 3',11.5)]:
        st=doc.styles[name];st.font.name=FONT;st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
        st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),FONT)
        st.paragraph_format.space_after=Pt(6);st.paragraph_format.line_spacing=1.18
        if name.startswith('Heading'):st.font.bold=True;st.paragraph_format.space_before=Pt(7)
    header=sec.header.paragraphs[0];font(header.add_run(AUTHOR+'  |  ACT仿真训练复盘'),8.5)
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT;font(footer.add_run('2026-10-24  ·  '),8.5)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    doc.core_properties.author=AUTHOR;doc.core_properties.title='ACT 抓放仿真训练复现与调参复盘'
    lines=source.read_text('utf-8').splitlines(); i=0; pending_break=False
    while i<len(lines):
        line=lines[i]
        if not line.strip():i+=1;continue
        if line=='<!-- page -->':pending_break=True;i+=1;continue
        if line.startswith('|'):
            block=[]
            while i<len(lines) and lines[i].startswith('|'):block.append(lines[i]);i+=1
            make_table(doc,block);continue
        if line.startswith('```'):
            code=[];i+=1
            while i<len(lines) and not lines[i].startswith('```'):code.append(lines[i]);i+=1
            p=doc.add_paragraph();p.paragraph_format.line_spacing=1.08;p.paragraph_format.space_after=Pt(7)
            font(p.add_run('\n'.join(code)),8.8,name='Consolas');i+=1;continue
        image=re.fullmatch(r'!\[([^\]]+)\]\(([^)]+)\)',line)
        if image:
            p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(output/image.group(2)),width=Cm(12.4 if image.group(2)=='positions.png' else 15.4))
            p.paragraph_format.space_after=Pt(2)
            caption=doc.add_paragraph();caption.alignment=WD_ALIGN_PARAGRAPH.CENTER
            font(caption.add_run('图  '+image.group(1)+'  ·  根据本项目保存数据绘制'),8.5);i+=1;continue
        if line.startswith('#'):
            marks,text=line.split(' ',1);level=0 if len(marks)==1 else len(marks)-1
            p=doc.add_heading(text,level=level)
            if pending_break:p.paragraph_format.page_break_before=True;pending_break=False
        else:
            is_list=bool(re.match(r'^\d+\. |^- ',line))
            p=doc.add_paragraph(style='List Bullet' if line.startswith('- ') else 'Normal')
            p.paragraph_format.space_after=Pt(5);p.paragraph_format.line_spacing=1.18
            runs(p,line[2:] if line.startswith('- ') else line,10 if is_list or line.startswith('实现：') or line.startswith('代码起点') else 11)
        i+=1
    path=output/'ACT仿真训练复现与调参复盘.docx';doc.save(path)
    record={'author':AUTHOR,'docx':str(path),'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),
            'source_md_sha256':hashlib.sha256(source.read_bytes()).hexdigest(),'plot_data_sha256':hashlib.sha256(dataset.read_bytes()).hexdigest(),
            'video_links':0,'embedded_video_objects':0,'new_training_steps':0,'new_physics_steps':0,'cpu_budget':CPU_BUDGET}
    (output/'build_record.json').write_text(json.dumps(record,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps(record,ensure_ascii=False))


if __name__=='__main__':main()
